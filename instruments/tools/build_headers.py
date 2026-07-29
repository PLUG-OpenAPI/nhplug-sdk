"""종목마스터 통합명세서(엑셀) → 상품별 C 헤더(.h) 생성기 (관리자용).

`.h` 는 **사람이 읽는 문서이자 파서의 입력(정본)** 이다.
파서(master.py)가 이 파일을 읽어 오프셋·길이를 얻으므로, 문서와 코드가 어긋날 수 없다.

사용:
    python instruments/tools/build_headers.py 종목마스터_통합명세서.xlsx

주의: 이 스크립트는 명세서가 갱신됐을 때만 실행하는 관리자용 도구다.
      일반 사용자는 생성된 instruments/headers/*.h 와 master.py 만 쓰면 된다.
"""
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("openpyxl 이 필요합니다:  pip install openpyxl")

SKIP_SHEETS = {"전체목록", "공통규칙·이슈"}
OUT_DIR = Path(__file__).resolve().parents[1] / "headers"


def _norm(v) -> str:
    return "" if v is None else str(v).strip()


def _clean(s: str) -> str:
    """주석에 들어갈 수 있게 개행·주석종료문자 제거."""
    return re.sub(r"\s+", " ", s.replace("*/", "* /")).strip()


def parse_index(wb):
    """전체목록 시트 → {파일명stem: {url, size, group, note}}"""
    ws = wb["전체목록"]
    rows = [[_norm(c) for c in r] for r in ws.iter_rows(values_only=True)]
    hdr = next(i for i, r in enumerate(rows) if r and r[0] == "No")
    cols = rows[hdr]
    idx = {}

    def col(keyword):
        return next((j for j, c in enumerate(cols) if keyword in c), None)

    c_group, c_file, c_url = col("구분"), col("파일명"), col("다운로드")
    c_size, c_svc, c_note = col("레코드"), col("관련서비스"), col("비고")

    for r in rows[hdr + 1:]:
        if not r or not r[0].isdigit():
            continue
        fname = r[c_file]
        if not fname:
            continue
        size = re.sub(r"[^0-9]", "", r[c_size]) if c_size is not None else ""
        idx[Path(fname).stem] = {
            "file": fname,
            "url": r[c_url] if c_url is not None else "",
            "size": int(size) if size else 0,
            "group": r[c_group] if c_group is not None else "",
            "service": r[c_svc] if c_svc is not None else "",
            "note": r[c_note] if c_note is not None else "",
        }
    return idx


def parse_sheet(ws):
    """상품 시트 → (필드리스트, 선언 레코드크기, 구조체명)"""
    rows = [[_norm(c) for c in r] for r in ws.iter_rows(values_only=True)]
    hdr = next(i for i, r in enumerate(rows) if r and r[0] == "No" and any("오프셋" in c for c in r))
    cols = rows[hdr]

    def col(kw):
        return next((j for j, c in enumerate(cols) if kw in c), None)

    c_name, c_ledger = col("항목명"), col("원장 필드명")
    c_off, c_len, c_desc, c_code = col("오프셋"), col("항목길이"), col("설명"), col("코드값")

    # 구조체명: 헤더행 위쪽 '구조체명' 라벨 다음 행에서 추출
    struct = ws.title
    for i, r in enumerate(rows[:hdr]):
        if any("구조체명" in c for c in r) and i + 1 < len(rows):
            cand = [c for c in rows[i + 1] if c]
            if len(cand) >= 2:
                struct = cand[1]
            break

    fields, declared = [], None
    for r in rows[hdr + 1:]:
        if not r or not r[0]:
            continue
        if "전체 행 사이즈" in r[0]:
            nums = [c for c in r if c.isdigit()]
            declared = int(nums[-1]) if nums else None
            break
        try:
            length = int(r[c_len])
            offset = int(r[c_off])
        except (ValueError, TypeError, IndexError):
            continue
        fields.append({
            "name": r[c_name] or f"f{offset}",
            "ledger": r[c_ledger] if c_ledger is not None else "",
            "offset": offset,
            "length": length,
            "desc": r[c_desc] if c_desc is not None else "",
            "code": r[c_code] if c_code is not None else "",
        })
    return fields, declared, struct


def c_identifier(name: str) -> str:
    """헤더에 쓸 안전한 C 식별자로 정규화(한글·괄호 등 제거)."""
    s = re.sub(r"[^0-9A-Za-z_]", "_", name).strip("_")
    if not s or s[0].isdigit():
        s = "f_" + s
    return s


def build_header(stem, meta, fields, declared, struct) -> str:
    L = []
    L.append("/" + "*" * 76)
    L.append(f" * NH투자증권 종목마스터 — {meta.get('group','')}".rstrip())
    L.append(" *")
    L.append(f" * @file      {meta.get('file', stem + '.mst')}")
    L.append(f" * @url       {meta.get('url','')}")
    L.append(f" * @record    {declared}      // 레코드 크기(byte). 파일크기 % record == 0 이어야 함")
    L.append(" * @encoding  CP949")
    L.append(" * @padding   공백(0x20) 우측 · 레코드 끝 1바이트 LF(0x0A) · 파일 헤더 없음")
    if meta.get("service"):
        L.append(f" * @service   {meta['service']}")
    if meta.get("note"):
        L.append(f" * @note      {_clean(meta['note'])}")
    L.append(" *")
    L.append(" * 파싱 주의: 반드시 'rb'(바이너리) 로 열 것. NUL 종료 문자열이 아니므로")
    L.append(" *           strlen 금지 — 길이 기반 슬라이싱 후 우측 공백 제거.")
    L.append(" " + "*" * 75 + "/")
    L.append("")
    L.append("#pragma pack(1)")
    L.append("typedef struct")
    L.append("{")
    for f in fields:
        decl = f"    char {c_identifier(f['name'])}[{f['length']}];"
        desc = _clean(f["desc"])
        note = f"@{f['offset']:<4d} {desc}"
        # 설명에 이미 <원장필드명> 이 들어 있으면 중복 표기하지 않는다
        ledger = _clean(f["ledger"])
        if ledger and ledger != f["name"] and f"<{ledger}>" not in desc:
            note += f" <{ledger}>"
        if f["code"]:
            note += f" | {_clean(f['code'])}"
        L.append(f"{decl:<44s}/* {note} */")
    L.append(f"}}   {c_identifier(struct).upper()};   /* sizeof = {declared} */")
    L.append("")
    return "\n".join(L)


def main():
    if len(sys.argv) < 2:
        raise SystemExit("사용: python build_headers.py <종목마스터_통합명세서.xlsx>")
    xlsx = Path(sys.argv[1])
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    index = parse_index(wb)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    made, bad = 0, []
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        fields, declared, struct = parse_sheet(ws)
        meta = index.get(ws.title, {})
        total = sum(f["length"] for f in fields)
        if declared is None or total != declared:
            bad.append(f"{ws.title}: 길이합 {total} != 선언 {declared}")
            continue
        if meta.get("size") and meta["size"] != declared:
            bad.append(f"{ws.title}: 목록 {meta['size']} != 시트 {declared}")
            continue
        (OUT_DIR / f"{ws.title}.h").write_text(
            build_header(ws.title, meta, fields, declared, struct), encoding="utf-8")
        made += 1
        print(f"  ✓ {ws.title}.h  ({len(fields):2d} fields, {declared} B)")

    print(f"\n생성 {made}개 → {OUT_DIR}")
    if bad:
        print("⚠️ 건너뜀:")
        for b in bad:
            print("   " + b)


if __name__ == "__main__":
    main()
